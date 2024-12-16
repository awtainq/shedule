import openpyxl
import datetime
import re

wb = openpyxl.load_workbook('Book2.xlsx')

def delete_dup(l):
    n = []
    for i in l:
        if i not in n:
            n.append(i)
    return n

class_groupings = {}            
my_list = []    
for sheet in wb:
    if 'Классы' in sheet.title:
        sheet = wb[sheet.title]
        break
    stop = 3
for i in range(2, sheet.max_row):
    my_list = []
    for cell in sheet[i][:11]:
        my_list.append(cell.value)
    if set(my_list) == {None}:
        stop -= 1
        if stop == 0:
            break
    else:
        if my_list[7] is not None:
            class_groupings[my_list[1]] = my_list[7]
for class_part in class_groupings:
    string = class_groupings[class_part].split(',')
    main = [string[0]]
    mateng = []
    eng = []
    inf = []
    pe = []
    hf = []
    Hf = []
    mms = []
    for part in string:
        part = part.strip()
        if 'Англ/Мат' in part:
            mateng.append(part)
        elif 'Анг' in part:
            eng.append(part)
        elif 'инф' in part.lower() or 'икт' in part.lower():
            inf.append(part)
        elif 'Физ-ра' in part:
            pe.append(part)
        elif 'Химбиочасть' in part or 'Физматчасть' in part:
            hf.append(part)
        elif 'Физчасть' in part or 'Хбчасть' in part:
            Hf.append(part)
        elif 'мат/физ/хим' in part.lower():
            mms.append(part)
    arr = [main, mateng, eng, inf, pe, hf, Hf, mms]
    arr = [x for x in arr if x]
    class_groupings[class_part] = arr

my_list = []
lesson_rooms = {} 
for sheet in wb:
    if 'Предметы' in sheet.title:
        sheet = wb[sheet.title]
        break
    stop = 3
for i in range(2, sheet.max_row):
    my_list = []
    for cell in sheet[i][:6]:
        my_list.append(cell.value)
    if set(my_list) == {None}:
        stop -= 1
        if stop == 0:
            break
    else:
        if my_list[3] != '' and my_list[3] != None:
            lesson_rooms[my_list[2]] = my_list[3].split(', ')

my_list = []
teachers_rooms = {} 
for sheet in wb:
    if 'Список учителей' in sheet.title:
        sheet = wb[sheet.title]
        break
    stop = 3
for i in range(2, sheet.max_row):
    my_list = []
    for cell in sheet[i][:6]:
        my_list.append(cell.value)
    if set(my_list) == {None}:
        stop -= 1
        if stop == 0:
            break
    else:
        if my_list[2] != '' and int(my_list[2]) > 0:
            teachers_rooms[my_list[1]] = my_list[5].split(', ')

data = []
for sheet in wb:
    if 'Нагрузка' in sheet.title:
        sheet = wb[sheet.title]
        break
    stop = 3
    teacher = ''
    
for i in range(2, sheet.max_row):
    my_list = []
    for cell in sheet[i][:11]:
        my_list.append(cell.value)
    if set(my_list) == {None}:
        stop -= 1
        if stop == 0:
            break
    else:
        stop = 3
        if my_list[0] is not None:
            teacher = my_list[0]
        else:
            if isinstance(my_list[1], datetime.datetime):
                class_name = (my_list[1].strftime('%-m-%-d'))
            else:
                class_name = my_list[1]
            class_part = my_list[2]
            lesson = my_list[3]
            count = my_list[5]
            if lesson in lesson_rooms:
                rooms = lesson_rooms[lesson]
            else:
                rooms = delete_dup(teachers_rooms[teacher] + str(my_list[9]).split(', '))
            data.append([teacher, class_name, class_part, lesson, count, rooms])

def sanitize(name):
    return re.sub(r'\W+', '_', str(name))

data = delete_dup(data)

# Очистка данных
sanitized_data = []
for entry in data:
    teacher, class_name, class_part, lesson, sessions_per_week, rooms = entry
    teacher = sanitize(teacher)
    class_name = sanitize(class_name)
    class_part = sanitize(class_part)
    lesson = sanitize(lesson)
    sessions_per_week = int(sessions_per_week)
    rooms = [sanitize(room) for room in rooms]
    sanitized_data.append([teacher, class_name, class_part, lesson, sessions_per_week, rooms])
data = sanitized_data

sanitized_class_groupings = {}
for class_name in class_groupings:
    sanitized_class_groupings[sanitize(class_name)] = [[sanitize(part) for part in parts] for parts in class_groupings[class_name]]
class_groupings = sanitized_class_groupings
  
def ret():
    return data, class_groupings

from collections import Counter
from itertools import chain
dats = []
for dat in data[:250]:
    dats.append(dat[5])
    print(dat)
c = Counter(list(chain.from_iterable(dats)))
print(c)